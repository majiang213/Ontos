# -*- coding: utf-8 -*-
"""Post-process md2docx output for mvp_plan.md: CJK fonts, headings, subtitle,
native-styled tables (gray bold header, thin gray borders, repeat header),
Consolas code block, page numbers."""
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = '/mnt/agents/output/md2docx_work/mvp_plan.footnote.docx'
DST = '/mnt/agents/output/Ontology平台MVP开发方案.docx'

SONG = u'宋体'
HEI = u'黑体'
LATIN = 'Times New Roman'
MONO = 'Consolas'
DARK = RGBColor(0x1F, 0x1F, 0x1F)
GRAY = RGBColor(0x59, 0x59, 0x59)

doc = docx.Document(SRC)
styles = {s.name: s for s in doc.styles}


def set_style_font(name, latin, eastasia, size_pt, bold=False, color=None):
    st = styles[name]
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), latin)
    rfonts.set(qn('w:hAnsi'), latin)
    rfonts.set(qn('w:cs'), latin)
    rfonts.set(qn('w:eastAsia'), eastasia)
    st.font.size = Pt(size_pt)
    st.font.bold = bold
    if color is not None:
        st.font.color.rgb = color
    return st


def set_spacing(st, before=None, after=None, line=None):
    pf = st.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


# ---------- 1. Base styles ----------
n = set_style_font('Normal', LATIN, SONG, 12, color=DARK)
set_spacing(n, line=1.4)

for nm in ('Body Text', 'First Paragraph'):
    st = set_style_font(nm, LATIN, SONG, 12, color=DARK)
    set_spacing(st, before=0, after=6, line=1.4)
    ppr = st.element.get_or_add_pPr()
    ind = ppr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        ppr.append(ind)
    ind.set(qn('w:firstLineChars'), '200')
    ind.set(qn('w:firstLine'), '0')

st = set_style_font('Compact', LATIN, SONG, 10.5, color=DARK)
set_spacing(st, before=2, after=2, line=1.3)

# ---------- 2. Headings (黑体) ----------
t = set_style_font('Title', LATIN, HEI, 22, bold=True, color=DARK)
t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(t, before=12, after=10)

h1 = set_style_font('Heading 1', LATIN, HEI, 20, bold=True, color=DARK)
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(h1, before=12, after=10)

h2 = set_style_font('Heading 2', LATIN, HEI, 15, bold=True, color=DARK)
set_spacing(h2, before=16, after=8)

h3 = set_style_font('Heading 3', LATIN, HEI, 13, bold=True, color=DARK)
set_spacing(h3, before=10, after=5)

h4 = set_style_font('Heading 4', LATIN, HEI, 12, bold=True, color=DARK)
set_spacing(h4, before=8, after=4)

sub = set_style_font('Subtitle', LATIN, SONG, 11.5, color=GRAY)
sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(sub, before=0, after=18)

bt = set_style_font('Block Text', LATIN, SONG, 10.5, color=GRAY)
set_spacing(bt, before=4, after=4, line=1.3)

# Source Code: monospace, keep indentation, no first-line indent
sc = set_style_font('Source Code', MONO, SONG, 9.5, color=DARK)
set_spacing(sc, before=6, after=6, line=1.25)
ppr = sc.element.get_or_add_pPr()
ind = ppr.find(qn('w:ind'))
if ind is None:
    ind = OxmlElement('w:ind')
    ppr.append(ind)
ind.set(qn('w:left'), '284')  # slight block indent, preserve inner spaces
shd = OxmlElement('w:shd')
shd.set(qn('w:val'), 'clear')
shd.set(qn('w:fill'), 'F2F2F2')
ppr.append(shd)

# ---------- 3. Title & subtitle paragraphs ----------
paras = doc.paragraphs
for p in paras:  # first (only) H1 -> Title
    if p.style.name == 'Heading 1':
        p.style = styles['Title']
        break
for p in paras:  # leading blockquote (版本...) -> Subtitle
    if p.style.name == 'Block Text' and u'版本' in p.text:
        p.style = styles['Subtitle']

# ---------- 4. Run-level font enforcement ----------
def fix_run_fonts(paragraph, eastasia, latin=None):
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.insert(0, rfonts)
        rfonts.set(qn('w:eastAsia'), eastasia)
        if latin:
            rfonts.set(qn('w:ascii'), latin)
            rfonts.set(qn('w:hAnsi'), latin)

HEAD_STYLES = {'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'}
for p in doc.paragraphs:
    if p.style.name in HEAD_STYLES:
        fix_run_fonts(p, HEI, LATIN)
    elif p.style.name == 'Source Code':
        fix_run_fonts(p, SONG, MONO)
    else:
        fix_run_fonts(p, SONG, LATIN)

# ---------- 5. Tables: gray bold header, thin gray borders, repeat header ----------
def style_table(tbl):
    tblPr = tbl._tbl.tblPr
    borders = tblPr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for tag in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        e = borders.find(qn('w:%s' % tag))
        if e is None:
            e = OxmlElement('w:%s' % tag)
            borders.append(e)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), 'A6A6A6')  # thin gray, no color
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    if len(tbl.rows) > 0:
        hdr = tbl.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        if trPr.find(qn('w:tblHeader')) is None:
            trPr.append(OxmlElement('w:tblHeader'))  # repeat across pages
        for cell in hdr.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:fill'), 'D9D9D9')  # light gray header
            tcPr.append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            va = OxmlElement('w:vAlign')
            va.set(qn('w:val'), 'center')
            tcPr.append(va)
            for p in cell.paragraphs:
                for run in p.runs:
                    rpr = run._r.get_or_add_rPr()
                    rfonts = rpr.find(qn('w:rFonts'))
                    if rfonts is None:
                        rfonts = OxmlElement('w:rFonts')
                        rpr.insert(0, rfonts)
                    rfonts.set(qn('w:eastAsia'), SONG)
                    rfonts.set(qn('w:ascii'), LATIN)
                    rfonts.set(qn('w:hAnsi'), LATIN)
                    if run.font.size is None:
                        run.font.size = Pt(10.5)

for tbl in doc.tables:
    style_table(tbl)

# ---------- 6. Page setup + footer page number ----------
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.6)
sec.top_margin = sec.bottom_margin = Cm(2.54)
footer = sec.footer
fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
fp.text = ''
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
run.font.size = Pt(9)
run.font.color.rgb = GRAY

# ---------- 7. Core properties ----------
doc.core_properties.title = u'通用 Ontology 平台 MVP 开发方案（V2.2）'
doc.core_properties.author = u'Ontology 平台设计组'

doc.save(DST)
print('saved:', DST)

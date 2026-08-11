# -*- coding: utf-8 -*-
"""Post-process pandoc base docx: CJK fonts, heading hierarchy, table styling,
subtitle, captions, page numbers."""
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = '/mnt/agents/output/docx_assets/ontology_base.docx'
DST = '/mnt/agents/output/通用Ontology平台产品与技术架构设计.docx'

SONG = u'宋体'      # SimSun - body
HEI = u'黑体'       # SimHei - headings
KAI = u'楷体'       # KaiTi - subtitle/captions accents
LATIN_BODY = 'Times New Roman'
LATIN_HEAD = 'Arial'
DARK = RGBColor(0x1F, 0x1F, 0x1F)
GRAY = RGBColor(0x59, 0x59, 0x59)

doc = docx.Document(SRC)
_styles_raw = doc.styles
styles = {s.name: s for s in _styles_raw}  # pandoc styles lack UI-name mapping


def set_rfonts(rpr_owner_font, style_el, latin, eastasia):
    """Set fonts on a style element (ascii/hAnsi/eastAsia)."""
    rpr = style_el.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), latin)
    rfonts.set(qn('w:hAnsi'), latin)
    rfonts.set(qn('w:cs'), latin)
    rfonts.set(qn('w:eastAsia'), eastasia)


def style_font(name, latin, eastasia, size_pt, bold=False, color=None,
               italic=False):
    st = styles[name]
    el = st.element
    set_rfonts(None, el, latin, eastasia)
    st.font.size = Pt(size_pt)
    st.font.bold = bold
    st.font.italic = italic
    if color is not None:
        st.font.color.rgb = color
    return st


def set_spacing(st, before=None, after=None, line=None):
    pf = st.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


def set_first_line_indent_chars(st, chars=200):
    ppr = st.element.get_or_add_pPr()
    ind = ppr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        ppr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), '0')


# ---------- 1. Base styles ----------
n = style_font('Normal', LATIN_BODY, SONG, 12, color=DARK)
set_spacing(n, line=1.4)

for nm in ('Body Text', 'First Paragraph'):
    st = style_font(nm, LATIN_BODY, SONG, 12, color=DARK)
    set_spacing(st, before=0, after=6, line=1.4)
    set_first_line_indent_chars(st, 200)

# Compact (table cells / tight lists): no indent, smaller
st = style_font('Compact', LATIN_BODY, SONG, 10.5, color=DARK)
set_spacing(st, before=1, after=1, line=1.15)

# ---------- 2. Headings ----------
h1 = style_font('Heading 1', LATIN_HEAD, HEI, 22, bold=True, color=DARK)
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(h1, before=12, after=12)

h2 = style_font('Heading 2', LATIN_HEAD, HEI, 16, bold=True, color=DARK)
set_spacing(h2, before=18, after=10)
h2.paragraph_format.page_break_before = True  # each chapter on new page

h3 = style_font('Heading 3', LATIN_HEAD, HEI, 14, bold=True, color=DARK)
set_spacing(h3, before=12, after=6)

h4 = style_font('Heading 4', LATIN_HEAD, HEI, 12, bold=True, color=DARK)
set_spacing(h4, before=8, after=4)

# Title / Subtitle / captions / quote
t = style_font('Title', LATIN_HEAD, HEI, 24, bold=True, color=DARK)
t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(t, before=36, after=18)

sub = style_font('Subtitle', LATIN_BODY, KAI, 14, color=GRAY)
sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(sub, before=0, after=24)

cap = style_font('Image Caption', LATIN_BODY, KAI, 9, color=GRAY)
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(cap, before=2, after=10)

bt = style_font('Block Text', LATIN_BODY, SONG, 10.5, color=GRAY)
set_spacing(bt, before=4, after=4, line=1.3)

# ---------- 3. Title & subtitle paragraphs ----------
paras = doc.paragraphs
# first Heading 1 -> Title
for p in paras:
    if p.style.name == 'Heading 1':
        p.style = styles['Title']
        break
# Block Text paragraph with 版本 -> Subtitle
for p in paras:
    if p.style.name == 'Block Text' and u'版本' in p.text:
        p.style = styles['Subtitle']

# ---------- 4. Run-level eastAsia font enforcement ----------
def fix_run_fonts(paragraph, eastasia, latin=None):
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.insert(0, rfonts)
        rfonts.set(qn('w:eastAsia'), eastasia)
        if latin:
            rfonts.set(qn('w:ascii'), latin)
            rfonts.set(qn('w:hAnsi'), latin)

HEAD_STYLES = {'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'}
for p in doc.paragraphs:
    if p.style.name in HEAD_STYLES:
        fix_run_fonts(p, HEI, LATIN_HEAD)
    else:
        fix_run_fonts(p, SONG)

# ---------- 5. Tables: light-gray header + thin borders ----------
def style_table(tbl):
    tblPr = tbl._tbl.tblPr
    # borders
    borders = tblPr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for tag, sz in (('top', 8), ('bottom', 8), ('left', 4), ('right', 4),
                    ('insideH', 4), ('insideV', 4)):
        e = borders.find(qn('w:%s' % tag))
        if e is None:
            e = OxmlElement('w:%s' % tag)
            borders.append(e)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '404040' if sz == 8 else '808080')
    # width 100%
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    # header row: shading + bold + repeat
    if len(tbl.rows) > 0:
        hdr = tbl.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader')
        trPr.append(th)
        for cell in hdr.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:fill'), 'D9D9D9')
            tcPr.append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
    # all cells: fonts + vertical center
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            va = OxmlElement('w:vAlign')
            va.set(qn('w:val'), 'center')
            tcPr.append(va)
            for p in cell.paragraphs:
                for run in p.runs:
                    rpr = run._r.get_or_add_rPr()
                    rfonts = rpr.find(qn('w:rFonts'))
                    if rfonts is None:
                        rfonts = OxmlElement('w:rFonts')
                        rpr.insert(0, rfonts)
                    rfonts.set(qn('w:eastAsia'), SONG)
                    rfonts.set(qn('w:ascii'), LATIN_BODY)
                    rfonts.set(qn('w:hAnsi'), LATIN_BODY)
                    if run.font.size is None:
                        run.font.size = Pt(10.5)

for tbl in doc.tables:
    style_table(tbl)

# ---------- 6. Footer with centered page number ----------
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.6)
sec.top_margin = sec.bottom_margin = Cm(2.54)
footer = sec.footer
fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
fp.text = ''
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
run.font.size = Pt(9)
run.font.color.rgb = GRAY

# ---------- 7. Core properties ----------
doc.core_properties.title = u'通用 Ontology 平台：产品与技术架构设计'
doc.core_properties.author = u'Ontology 平台设计组'

doc.save(DST)
print('saved:', DST)
